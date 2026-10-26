"""Build the presentation assets from release evidence and the timed script.

This documentation utility never imports or executes the compiler. All slide
text/charts remain editable in PowerPoint; PDF and PNG use embedded fonts.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from xml.sax.saxutils import escape

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
WIDTH, HEIGHT = 1920, 1080
BG = '#0C1422'
CARD = '#142136'
LINE = '#2C3C52'
INK = '#F1F5FA'
MUTED = '#A7B8CF'
MINT = '#5CE4BF'
BLUE = '#95B4FF'
GOLD = '#F1CC83'
FONT_FILES = {
    'regular': Path('/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf'),
    'bold': Path('/usr/share/fonts/dejavu-sans-fonts/DejaVuSans-Bold.ttf'),
    'mono': Path('/usr/share/fonts/liberation-mono/LiberationMono-Regular.ttf'),
}
PPT_FONTS = {'regular': 'DejaVu Sans', 'bold': 'DejaVu Sans', 'mono': 'Liberation Mono'}
FONT_CACHE = {}


def read(path):
    return json.loads(path.read_text())


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font(size, face='regular'):
    key = size, face
    if key not in FONT_CACHE:
        FONT_CACHE[key] = ImageFont.truetype(str(FONT_FILES[face]), round(size))
    return FONT_CACHE[key]


def text_width(value, size, face='regular'):
    return font(size, face).getlength(value)


class Slide:
    def __init__(self, number, title, *, total=9, reference=False, source=None):
        self.number, self.title, self.items = number, title, []
        self.reference = reference
        self.rect(0, 0, WIDTH, HEIGHT, BG)
        self.rect(96, 66, 48, 7, MINT)
        self.text('DEUTERON / PASS REFERENCE' if reference else 'DEUTERON / COMPILER PASSES', 164, 53, 25, MINT, 'bold', width=1250)
        self.text(f'{number:02d} / {total:02d}', 1665, 53, 25, MUTED, width=150, align='right')
        self.line(96, 1004, 1824, 1004, LINE, 2)
        self.text(source or ('Reference only · outside the three-minute video' if reference else 'Scheduled Emmy IR → NVIDIA PTX'), 100, 1023, 21, MUTED, width=1240)
        self.text('RTX 5090  /  SM120', 1370, 1023, 22, MUTED, width=450, align='right')

    def rect(self, x, y, w, h, fill, radius=0, stroke=None):
        self.items.append(dict(kind='rect', x=x, y=y, w=w, h=h, fill=fill, radius=radius, stroke=stroke))

    def line(self, x, y, x2, y2, color, width=2):
        self.items.append(dict(kind='line', x=x, y=y, x2=x2, y2=y2, color=color, width=width))

    def text(self, value, x, y, size, color=INK, face='regular', *, width=1700, align='left', leading=1.2):
        lines = value.split('\n')
        for line in lines:
            actual = text_width(line, size, face)
            assert actual <= width, (self.number, line, actual, width)
        height = len(lines) * size * leading + 8
        assert x >= 0 and y >= 0 and x + width <= WIDTH and y + height <= HEIGHT, (self.number, value)
        self.items.append(dict(kind='text', value=value, x=x, y=y, size=size, color=color,
                               face=face, w=width, h=height, align=align, leading=leading))

    def wrapped(self, value, x, y, size, color=INK, face='regular', *, width=500, leading=1.25):
        lines = []
        for paragraph in value.split('\n'):
            current = ''
            for word in paragraph.split():
                candidate = current + ' ' + word if current else word
                if current and text_width(candidate, size, face) > width:
                    lines.append(current)
                    current = word
                else:
                    current = candidate
            lines.append(current)
        self.text('\n'.join(lines), x, y, size, color, face, width=width, leading=leading)

    def arrow(self, x, y, width=38, color=MINT):
        self.line(x, y, x + width, y, color, 4)
        self.line(x + width - 11, y - 10, x + width, y, color, 4)
        self.line(x + width - 11, y + 10, x + width, y, color, 4)


def heading(s, title, subtitle):
    s.text(title, 108, 153, 66, INK, 'bold', width=1700)
    s.text(subtitle, 112, 263, 29, MUTED, width=1685)


def create_slides(warm):
    gm = warm['geometric_mean_latency_us']
    assert len(warm['calls']) == 270 and all(call['correct'] for call in warm['calls'])
    slides = []

    s = Slide(1, 'The compiler transformations we built')
    s.text('Deuteron', 110, 221, 150, INK, 'bold', width=1650)
    s.text('The compiler transformations we built', 120, 428, 48, INK, width=1670)
    s.text('Scheduled Emmy IR → PTX  ·  RTX 5090', 120, 518, 38, MUTED, width=1670)
    s.text('Recognize structure. Check the rewrite. Change the GPU work.', 120, 606, 32, MINT, width=1680)
    for x, title, body, accent in (
        (112, 'Reuse data', 'Vectorize loads\nKeep reduction inputs', MINT),
        (690, 'Redistribute work', 'Cooperate across lanes\nFold eligible K partitions', BLUE),
        (1268, 'Build pipelines', 'Pack tensor operands\nStage attention weights', GOLD),
    ):
        s.rect(x, 721, 540, 222, CARD, 24)
        s.rect(x, 721, 540, 5, accent)
        s.text(title, x + 30, 752, 37, accent, 'bold', width=480)
        s.text(body, x + 30, 819, 29, INK, width=480, leading=1.3)
    slides.append(s)

    s = Slide(2, 'Where our passes run', source='compiler/build.py · compiler/analysis.py · compiler/codegen/lower.py')
    heading(s, 'Where our passes run', 'Our contribution: structural recognizers, rewrite guards, and transformed GPU schedules.')
    nodes = [
        (110, 320, 'Validate', 'Scheduled IR\n+ launch metadata\nTypes and contract', MINT),
        (485, 470, 'Rewrite', 'Vectorize → forward loads\nReorder → fold split-K', MINT),
        (1010, 365, 'Lower', 'Reductions\nTensor pipelines\nStaged recurrences', BLUE),
        (1430, 380, 'LLVM → PTX', 'Typed LLVM IR\nLLVM optimization\nNVPTX code generation', GOLD),
    ]
    for i, (x, width, title, body, accent) in enumerate(nodes):
        s.rect(x, 365, width, 339, CARD, 24)
        s.text(f'0{i+1}', x + 27, 396, 28, accent, 'bold', width=90)
        s.text(title, x + 27, 452, 35, INK, 'bold', width=width-54)
        s.text(body, x + 27, 521, 27, MUTED, width=width-54, leading=1.55)
        if i < 3:
            s.arrow(x + width + 10, 528, 28)
    s.text('Actual early rewrite order', 114, 771, 28, MINT, 'bold', width=1680)
    s.text('vectorize.py  →  regforward.py  →  recurrence.py  →  gridfold.py', 114, 824, 30, INK, 'mono', width=1680)
    s.text('Each rewrite checks its own indices, ownership, and lifetime conditions.', 114, 915, 31, MUTED, width=1690)
    slides.append(s)

    s = Slide(3, 'Reuse data before loading it again', source='compiler/passes/vectorize.py · compiler/passes/regforward.py · compiler/transforms.py')
    heading(s, 'Reuse data before reloading it', 'Elementwise vectorization · Reduction load forwarding · Coalesced sum/square reduction')
    s.rect(110, 345, 810, 236, CARD, 23)
    s.rect(1000, 345, 810, 236, CARD, 23)
    s.text('Scalar loop', 142, 376, 34, MUTED, 'bold', width=740)
    for i in range(4):
        x = 145 + i * 181
        s.rect(x, 456, 148, 70, LINE, 10)
        s.text('load', x + 31, 474, 29, INK, 'mono', width=100)
        if i < 3:
            s.arrow(x + 153, 491, 20, MUTED)
    s.text('Four contiguous values per lane', 1032, 376, 34, MINT, 'bold', width=744)
    for i in range(4):
        x = 1041 + i * 173
        s.rect(x, 456, 165, 70, '#24594F', 8)
        s.text(f'x{i}', x + 61, 474, 30, MINT, 'mono', width=80)
    s.rect(110, 632, 810, 241, CARD, 23)
    s.text('Reduction load forwarding', 142, 661, 33, MINT, 'bold', width=743)
    s.text('load vector → reduce → consumer', 142, 737, 31, INK, width=743)
    s.text('Reuse the same registers; skip the reload.', 142, 801, 28, MUTED, width=743)
    s.rect(1000, 632, 810, 241, CARD, 23)
    s.text('Coalesced sum/square reduction', 1032, 661, 32, BLUE, 'bold', width=744)
    s.text('Contiguous load batches\nLong loops: independent accumulators', 1032, 733, 31, INK, width=744, leading=1.5)
    s.text('Pass-specific guards: matching indices · readonly inputs · closed lifetimes · complete coverage',
           114, 931, 27, MUTED, width=1690)
    slides.append(s)

    s = Slide(4, 'Reshape parallel work', source='compiler/transforms.py · compiler/passes/tailtree.py · compiler/passes/gridfold.py')
    heading(s, 'Reshape parallel work', 'Cooperative dot reduction · Small-tree elimination · Split-K grid folding')
    s.rect(110, 344, 810, 485, CARD, 24)
    s.rect(1000, 344, 810, 485, CARD, 24)
    s.text('Lane / warp work', 145, 380, 36, MINT, 'bold', width=740)
    s.text('Repeated dot → lane partials → shuffle', 145, 456, 31, INK, width=740)
    for i in range(4):
        x = 145 + 182 * i
        s.rect(x, 532, 156, 83, '#24594F', 10)
        s.text(f'p{i}', x + 53, 550, 37, MINT, 'mono', width=80)
    s.text('Closed small trees → direct sums', 145, 677, 33, INK, width=740)
    s.text('Warp 0 can combine up to 32 partials.', 145, 747, 27, MUTED, width=740)
    s.text('Split-K grid folding', 1034, 380, 36, BLUE, 'bold', width=740)
    s.text('Partitions per output tile in a larger grid', 1034, 441, 26, MUTED, width=740)
    for i in range(4):
        x = 1035 + i * 182
        s.rect(x, 495, 155, 71, '#2D4165' if i < 2 else '#1D2B40', 10)
        s.text(f'CTA {i}' if i < 2 else 'exit', x + 27, 515, 27,
               BLUE if i < 2 else MUTED, 'mono', width=115)
    s.text('Logical K partitions:   4 → 2\nK work / active CTA:  1× → 2×', 1034, 603, 33, INK, width=740, leading=1.55)
    s.text('Same physical grid. Same atomic output.', 1034, 747, 27, MUTED, width=740)
    s.text('Illustrative factor 2. Actual folding checks coverage, scratch capacity, K depth, and remaining parallelism.',
           114, 867, 25, MUTED, width=1690)
    s.text('Change who does the work while retaining the original launch contract.', 114, 936, 32, MINT, width=1690)
    slides.append(s)

    s = Slide(5, 'Turn contractions into tensor work', source='compiler/passes/contraction.py · compiler/codegen/tensor.py · compiler/storage.py')
    heading(s, 'Recover the contraction', 'Contraction tensorization · Tensor shared-memory layout')
    s.rect(110, 339, 725, 430, CARD, 24)
    s.text('Shared multiply + add', 145, 374, 36, INK, 'bold', width=656)
    s.text('Recognize a closed contraction\nProve affine tile coverage\nDerive full-warp ownership', 145, 459, 33, MUTED, width=650, leading=1.75)
    s.text('Guarded substitution → warp MMA', 145, 688, 31, MINT, width=650)
    s.arrow(858, 554, 49)
    s.rect(940, 339, 870, 430, CARD, 24)
    s.text('BF16x3 or TF32x3', 975, 374, 41, MINT, 'bold', width=798)
    s.text('A ≈ Ah + Al     B ≈ Bh + Bl', 975, 461, 30, MUTED, 'mono', width=796)
    s.text('Ah·Bh + Ah·Bl + Al·Bh', 975, 531, 43, INK, 'mono', width=796)
    s.text('FP32 buffers + accumulation', 975, 614, 34, INK, width=796)
    s.text('Low × low omitted → approximate products', 975, 688, 28, GOLD, width=796)
    s.text('Ordinary FP32 contraction mode is also available.', 114, 802, 27, MUTED, width=1690)
    s.rect(110, 858, 1700, 111, CARD, 20)
    s.text('SEPARATE FP16 STORAGE', 145, 883, 25, GOLD, 'bold', width=462)
    s.text('Pointwise storage and conversion', 145, 925, 22, MUTED, width=462)
    s.text('FP16 → explicit FP32 arithmetic → FP16', 650, 874, 35, INK, width=1120)
    s.text('Exact widening · nearest-even narrowing', 650, 927, 24, MUTED, width=1120)
    slides.append(s)

    s = Slide(6, 'Rewrite the pipeline feeding tensor cores', source='compiler/passes/regstage.py · compiler/codegen/tensor.py · compiler/codegen/lower.py')
    heading(s, 'Feed tensor cores, then drain', 'Packed register staging · Tensor fragment pipelining · Pipeline drain')
    s.rect(110, 341, 1700, 132, CARD, 22)
    s.text('SCHEDULED / NONPACKED', 141, 370, 25, BLUE, 'bold', width=440)
    s.text('FP32 cp.async ring → shared FP32 → convert → MMA', 603, 382, 35, INK, width=1170)
    s.rect(110, 515, 1700, 281, CARD, 24)
    s.text('Certified packed BF16 path', 144, 545, 35, MINT, 'bold', width=1600)
    for i, text_value in enumerate(('Vector global\nprefetch', 'Pack high +\nresidual', 'One used\nshared stage', 'Load next\nfragment', 'Current\nMMA')):
        x = 143 + 334 * i
        s.rect(x, 620, 286, 110, '#24594F', 15)
        s.text(text_value, x + 22, 638, 29, INK, width=245, leading=1.32)
        if i < 4:
            s.arrow(x + 296, 674, 26)
    s.text('Tile dataflow with global and fragment lookahead.', 145, 751, 24, MUTED, width=1590)
    s.text('Innovation: certify the exact copy phase and the packed producer / consumer layout.', 114, 836, 30, MINT, width=1690)
    s.text('Final tile: omit unused prefetch + consumer barrier; keep producer writes and their barrier.', 114, 893, 27, INK, width=1690)
    s.text('Eligible FP32 / TF32x3 paths retain 4 / 8 / 16-byte cp.async groups.', 114, 950, 25, MUTED, width=1690)
    slides.append(s)

    s = Slide(7, 'Recognize and stage the attention recurrence', source='compiler/passes/recurrence.py · compiler/passes/lifetime.py')
    heading(s, 'Stage the attention recurrence', 'Staged online recurrence · Blocked recurrence loads · Recurrence shared-storage reuse')
    s.rect(110, 340, 1700, 110, CARD, 20)
    s.text('BEFORE', 143, 368, 25, MUTED, 'bold', width=175)
    s.text('Output columns repeat score and online max / mass / value work', 352, 371, 33, INK, width=1400)
    stages = [('Scores once', 'per sharing group'), ('Maximum', 'across the keys'), ('Weights', 'exp(score − max)')]
    for i, (title, detail) in enumerate(stages):
        x = 110 + i * 430
        s.rect(x, 510, 370, 162, CARD, 22)
        s.text(title, x + 27, 540, 36, MINT, 'bold', width=316)
        s.text(detail, x + 27, 605, 27, INK, width=316)
        if i < 2:
            s.arrow(x + 386, 591, 27)
    # One computed weight branches into register mass and shared storage for PV.
    s.line(1348, 591, 1390, 591, MINT, 4)
    s.line(1390, 544, 1390, 651, MINT, 4)
    for y, title, accent in ((504, 'Register mass', MINT), (611, 'Shared weights → PV', BLUE)):
        s.arrow(1390, y + 43, 26, accent)
        s.rect(1430, y, 380, 88, CARD, 17)
        s.text(title, 1450, y + 25, 27, accent, 'bold', width=342)
    s.text('Unnormalized weights; normalization stays in the output suffix.', 114, 704, 25, MUTED, width=1690)
    for x, title, body, accent in (
        (110, 'Blocked recurrence loads', '2 / 4 adjacent keys per lane\nVector loads and hoisted addresses', BLUE),
        (1000, 'Recurrence shared-storage reuse', 'Same warp + same component\nRead → maximum → overwrite', GOLD),
    ):
        s.rect(x, 767, 810, 180, CARD, 20)
        s.text(title, x + 30, 790, 31, accent, 'bold', width=750)
        s.text(body, x + 30, 852, 28, INK, width=750, leading=1.35)
    slides.append(s)

    s = Slide(8, 'Redistribute probability–value work', source='compiler/passes/recurrence.py · _grouped_value_axis · grouped_values')
    heading(s, 'Redistribute probability–value work', 'Pass: Warp-distributed vector PV reduction')
    s.rect(110, 341, 650, 544, CARD, 24)
    s.text('BEFORE', 145, 378, 25, MUTED, 'bold', width=570)
    s.text('1 lane', 145, 442, 74, MINT, 'bold', width=570)
    s.text('1 column × every key', 145, 552, 39, INK, width=570)
    s.text('sum_k p[k] * V[k,c]', 145, 637, 34, MUTED, 'mono', width=570)
    s.text('Guard: 1 warp · 32 columns\nAligned loads + closed suffix', 145, 754, 28, MUTED, width=570, leading=1.45)
    s.rect(820, 341, 990, 544, CARD, 24)
    s.text('AFTER · 4 groups × 8 lanes', 853, 378, 34, INK, 'bold', width=925)
    s.text('Four adjacent columns per lane', 998, 437, 25, MUTED, width=750)
    for row in range(4):
        y = 490 + row * 71
        s.text(f'Shard {row}', 853, y + 13, 27, BLUE, width=142)
        for column in range(8):
            x = 1002 + 94 * column
            s.rect(x, y, 84, 54, '#24594F' if row % 2 == 0 else '#2D4165', 8)
            label = f'{column*4}–{column*4+3}'
            s.text(label, x + 10, y + 16, 20, INK, 'mono', width=70)
    s.text('Shuffle offsets 16 + 8 → first 8 lanes write', 853, 801, 29, MINT, width=925)
    s.rect(110, 916, 1700, 64, '#24594F', 16)
    s.text('Correct suffix replay: restore the saved denominator before every column.', 141, 932, 31, INK, width=1635)
    slides.append(s)

    s = Slide(9, 'Measure the whole compiler on RTX 5090')
    heading(s, 'RTX 5090 · whole-compiler results', 'Geometric mean warm latency · 15 open workloads · lower is better')
    x0, max_width, maximum = 340, 775, 7.0
    s.text('µs', 1110, 340, 26, MUTED, width=80)
    for label, key, y, accent in (
        ('Deuteron', 'Deuteron', 418, MINT),
        ('NVCC 12.8', 'NVCC12.8.93', 562, '#8496B0'),
        ('Triton 3.6', 'Triton3.6.0', 706, BLUE),
    ):
        value = gm[key]
        length = value / maximum * max_width
        s.text(label, 120, y + 17, 31, INK, 'bold', width=208)
        s.rect(x0, y, length, 78, accent, 8)
        s.text(f'{value:.3f}', x0 + length + 17, y + 17, 32, INK, 'bold', width=126)
    s.line(x0, 833, x0 + max_width, 833, LINE, 2)
    for value in (0, 2, 4, 6):
        x = x0 + value / maximum * max_width
        s.line(x, 833, x, 845, MUTED, 2)
        s.text(str(value), x - 7, 858, 24, MUTED, width=50)
    for y, comparison, label in ((373, 'NVCC12.8.93', 'vs generated CUDA / NVCC'), (637, 'Triton3.6.0', 'vs tested Triton baseline')):
        s.rect(1305, y, 505, 222, CARD, 23)
        s.text(f'{gm[comparison] / gm["Deuteron"]:.2f}×', 1340, y + 25, 100, MINT, 'bold', width=437)
        s.wrapped(label, 1340, y + 146, 27, INK, width=437)
    s.text('270 balanced warm calls · Original inputs and numerical tolerances · All output checks pass', 114, 921, 27, INK, width=1690)
    s.text('20 warmups / 200 iterations · Limited Triton search · Compile / assembly / transport excluded', 114, 969, 23, MUTED, width=1690)
    slides.append(s)
    return slides


def create_reference_slides(inventory):
    entries = inventory['entries']
    assert len(entries) == 27
    titles = ('Indexing, reuse, and reduction', 'Tensor and recurrence mechanisms', 'Normalization and final lowering')
    slides = []
    for page in range(3):
        s = Slide(page + 1, titles[page], total=3, reference=True)
        heading(s, titles[page], f'Entries {page*9+1}–{page*9+9} of 27 · Passes, analyses, policies, and lowering mechanisms')
        for index, entry in enumerate(entries[page*9:page*9+9]):
            x, y = 110 + (index % 3) * 580, 337 + (index // 3) * 221
            s.rect(x, y, 538, 205, CARD, 20)
            s.text(entry['id'], x + 24, y + 16, 18, MINT, 'bold', width=90)
            s.wrapped(entry['name'], x + 24, y + 45, 26, INK, 'bold', width=490, leading=1.13)
            s.wrapped(entry['summary'], x + 24, y + 116, 22, MUTED, width=490, leading=1.22)
            s.text(entry['source'], x + 24, y + 176, 16, BLUE, 'mono', width=490)
        slides.append(s)
    return slides


def color(value):
    return colors.HexColor(value)


def draw_pdf_slide(pdf, slide):
    # PDF points are half the 1920×1080 scene coordinates.
    pdf.saveState()
    pdf.scale(0.5, 0.5)
    for item in slide.items:
        kind = item['kind']
        if kind == 'rect':
            pdf.setFillColor(color(item['fill']))
            pdf.setStrokeColor(color(item.get('stroke') or item['fill']))
            pdf.roundRect(item['x'], HEIGHT - item['y'] - item['h'], item['w'], item['h'],
                          item['radius'], fill=1, stroke=int(bool(item.get('stroke'))))
        elif kind == 'line':
            pdf.setStrokeColor(color(item['color']))
            pdf.setLineWidth(item['width'])
            pdf.line(item['x'], HEIGHT - item['y'], item['x2'], HEIGHT - item['y2'])
        else:
            pdf.setFont(item['face'], item['size'])
            pdf.setFillColor(color(item['color']))
            for index, line in enumerate(item['value'].split('\n')):
                y = HEIGHT - item['y'] - item['size'] * (0.93 + index * item['leading'])
                if item['align'] == 'right':
                    pdf.drawRightString(item['x'] + item['w'], y, line)
                else:
                    pdf.drawString(item['x'], y, line)
    pdf.restoreState()


def build_pptx(slides, script, destination, inventory):
    deck = Presentation()
    deck.slide_width = Pt(WIDTH / 2)
    deck.slide_height = Pt(HEIGHT / 2)
    deck.core_properties.title = 'Deuteron — compiler passes and innovations' if script else 'Deuteron — complete pass reference'
    deck.core_properties.subject = 'Implemented compiler transformations and verified RTX 5090 measurements'
    deck.core_properties.author = 'Deuteron'
    if script:
        assert len(slides) == len(script['segments'])
    lookup = {entry['id']: entry for entry in inventory['entries']}
    for scene_index, scene in enumerate(slides):
        segment = script['segments'][scene_index] if script else None
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for item in scene.items:
            if item['kind'] == 'line':
                from pptx.enum.shapes import MSO_CONNECTOR
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                    Pt(item['x'] / 2), Pt(item['y'] / 2), Pt(item['x2'] / 2), Pt(item['y2'] / 2))
                shape.line.color.rgb = RGBColor.from_string(item['color'][1:])
                shape.line.width = Pt(item['width'] / 2)
            elif item['kind'] == 'rect':
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if item['radius'] else MSO_SHAPE.RECTANGLE,
                    Pt(item['x'] / 2), Pt(item['y'] / 2), Pt(item['w'] / 2), Pt(item['h'] / 2))
                if item['radius']:
                    shape.adjustments[0] = min(0.5, item['radius'] / min(item['w'], item['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item['fill'][1:])
                shape.line.fill.background()
            else:
                shape = slide.shapes.add_textbox(Pt(item['x'] / 2), Pt(item['y'] / 2), Pt(item['w'] / 2), Pt(item['h'] / 2))
                frame = shape.text_frame
                frame.clear()
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.vertical_anchor = MSO_ANCHOR.TOP
                frame.word_wrap = False
                frame.auto_size = MSO_AUTO_SIZE.NONE
                for index, line in enumerate(item['value'].split('\n')):
                    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    if item['align'] == 'right':
                        from pptx.enum.text import PP_ALIGN
                        paragraph.alignment = PP_ALIGN.RIGHT
                    paragraph.space_before = paragraph.space_after = Pt(0)
                    paragraph.line_spacing = Pt(item['size'] * item['leading'] / 2)
                    run = paragraph.add_run()
                    run.text = line
                    run.font.name = PPT_FONTS[item['face']]
                    run.font.bold = item['face'] == 'bold'
                    run.font.size = Pt(item['size'] / 2)
                    run.font.color.rgb = RGBColor.from_string(item['color'][1:])
        if segment:
            details = '\n'.join(f"{pid}: {lookup[pid]['name']} | {lookup[pid]['source']} | Guard: {lookup[pid]['guard']}"
                                for pid in segment['pass_ids'])
            notes = (f"{timestamp(segment['start_seconds'])}–{timestamp(segment['end_seconds'])} | {segment['title']}\n\n"
                     f"{segment['spoken']}\n\nVISUAL: {segment['visual_cue']}\n\n{details}\n\n"
                     'Evidence: passes.json, pass-reference.md, innovations.md, and recorded round-6 measurements. '
                     'These are rehearsal cues; align the edit to the actual recording. '
                     'The aggregate benchmark is not an individual-pass ablation.')
        else:
            notes = 'REFERENCE ONLY — outside the three-minute narrated deck.\n\n' + '\n\n'.join(
                f"{entry['id']}. {entry['name']}\n{entry['description']}\nGuard: {entry['guard']}\nSource: {entry['source']}"
                for entry in inventory['entries'][scene_index*9:scene_index*9+9])
        slide.notes_slide.notes_text_frame.text = notes
    deck.save(destination)


def timestamp(seconds):
    return f'{seconds // 60}:{seconds % 60:02d}'


def make_script_documents(script, destination):
    title = 'Deuteron · 3-minute video script'
    word_count = sum(len(segment['spoken'].split()) for segment in script['segments'])
    opening = (f'{word_count} spoken words · {len(script["segments"])} visual scenes · finish speaking by 2:55, then hold for five seconds. '
               'Record one clean, continuous take with no background music. Read only the narration below; '
               'the timecodes and visual cues are for rehearsal.')
    styles = {
        'title': ParagraphStyle('title', fontName='bold', fontSize=23, leading=28, textColor=color(BG), spaceAfter=12),
        'intro': ParagraphStyle('intro', fontName='regular', fontSize=10.5, leading=15.5, textColor=colors.HexColor('#526275'), spaceAfter=16),
        'head': ParagraphStyle('head', fontName='bold', fontSize=13, leading=17, textColor=colors.HexColor('#14624F'), spaceBefore=13, spaceAfter=8),
        'body': ParagraphStyle('body', fontName='regular', fontSize=12.4, leading=19, textColor=color(BG), spaceAfter=7),
        'cue': ParagraphStyle('cue', fontName='regular', fontSize=9.5, leading=13.5, textColor=colors.HexColor('#637083'), spaceAfter=8),
    }
    flow = [Paragraph(escape(title), styles['title']), Paragraph(escape(opening), styles['intro'])]
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = DocInches(0.68)
    section.left_margin = section.right_margin = DocInches(0.8)
    normal = document.styles['Normal']
    normal.font.name = 'DejaVu Sans'
    normal.font.size = DocPt(13)
    normal.paragraph_format.line_spacing = 1.2
    document.add_heading(title, 0)
    document.add_paragraph(opening)
    brief_cues = [segment['brief_cue'] for segment in script['segments']]
    for index, segment in enumerate(script['segments']):
        if index in (3, 6):
            continuation = 'Deuteron · reductions and tensors' if index == 3 else 'Deuteron · attention and results'
            flow += [PageBreak(), Paragraph(escape(continuation), styles['title'])]
            document.add_page_break()
        heading = f"{timestamp(segment['start_seconds'])}–{timestamp(segment['end_seconds'])}  |  Slide {index + 1}: {segment['title']}"
        flow.append(Paragraph(escape(heading), styles['head']))
        flow.append(Paragraph(escape(segment['spoken']), styles['body']))
        flow.append(Paragraph('<b>Cue:</b> ' + escape(brief_cues[index]), styles['cue']))
        document.add_heading(heading, 2)
        document.add_paragraph(segment['spoken'])
        cue = document.add_paragraph('Cue: ' + brief_cues[index])
        for run in cue.runs:
            run.font.size = DocPt(10)
            run.font.color.rgb = DocColor.from_string('637083')
    closing = ('2:55–3:00: stop speaking; hold slide 9. Leave extra silence on the raw recording for editing. '
               'The full recording setup, pronunciation notes, optional cuts, and edit plan are in the companion Markdown files.')
    flow.append(Paragraph(escape(closing), styles['intro']))
    document.add_paragraph(closing)
    pdf = SimpleDocTemplate(str(destination / 'recording-script.pdf'), pagesize=A4,
                            leftMargin=46, rightMargin=46, topMargin=40, bottomMargin=40,
                            title=title, author='Deuteron')

    def footer(page, _):
        page.setFont('regular', 9)
        page.setFillColor(colors.HexColor('#637083'))
        page.drawString(46, 22, 'Deuteron · narration guide · verified release efb728d')
        page.drawRightString(A4[0] - 46, 22, str(page.getPageNumber()))

    pdf.build(flow, onFirstPage=footer, onLaterPages=footer)
    document.save(destination / 'recording-script.docx')
    (destination / 'narration.txt').write_text('\n\n'.join(segment['spoken'] for segment in script['segments']) + '\n')
    return word_count


def make_presentation_guide(script, inventory, destination):
    count = sum(len(segment['spoken'].split()) for segment in script['segments'])
    lookup = {entry['id']: entry for entry in inventory['entries']}
    lines = [
        '# Deuteron: compiler passes and innovations in three minutes', '',
        f'Read only the quoted narration aloud. The script has **{count} spoken words**, nine narrated slides, and a five-second silent hold after 2:55. The three separate pass-reference slides are for preparation and discussion; they are outside the three-minute edit.', '',
        'Record one continuous take with clear speech. The slides carry the diagrams and exact pass names; keep the narration conversational. [script.json](script.json) is the timing source. The [engineering contributions](innovations.md) and [complete pass reference](pass-reference.md) explain the implementation and guards.', '',
        '| Slide | Time | Words | Main point |', '|---|---|---:|---|',
    ]
    for segment in script['segments']:
        lines.append(f"| {segment['slide']} | {timestamp(segment['start_seconds'])}–{timestamp(segment['end_seconds'])} | {len(segment['spoken'].split())} | {segment['title']} |")
    lines += ['| 9, held | 2:55–3:00 | 0 | Silent results-slide hold |', '']
    for segment in script['segments']:
        lines += [f"## {timestamp(segment['start_seconds'])}–{timestamp(segment['end_seconds'])} · Slide {segment['slide']}: {segment['title']}", '',
                  '**Delivery:** ' + segment['delivery'], '', '> ' + segment['spoken'], '', '**On screen:** ' + segment['visual_cue'], '']
        if segment['pass_ids']:
            names = [f"{pid} · {lookup[pid]['name']}" for pid in segment['pass_ids']]
            lines += ['**Implemented mechanisms:** ' + '; '.join(names) + '.', '']
    lines += [
        '## What to call innovative', '',
        'Present the contribution as the compiler recognizers, legality checks, ownership rewrites, and composition of GPU transformations. The exact copy-phase proof, packed producer/consumer certificate, guarded attention storage reuse, and mutable-suffix replay are concrete engineering details. These use established GPU algorithms and instructions. LLVM supplies its optimization and NVPTX backend.', '',
        'The optional [three-page pass reference](pass-reference.pdf) covers all 27 documented entries. That count includes analyses, policies, and lowering mechanisms; it does not mean 27 independent pass modules. The [source-linked explanations](pass-reference.md) give the before/after behavior and guards.', '',
        '## Pronunciation', '',
        '| Written | Say |', '|---|---|', '| IR / PTX | “I–R” / “P–T–X” |', '| LLVM / NVPTX | “L–L–V–M” / “N–V–P–T–X” |', '| BF16x3 | “bee-eff sixteen times three” |', '| TF32x3 | “tee-eff thirty-two times three” |', '| FP16 / FP32 | “eff-pee sixteen” / “eff-pee thirty-two” |', '| MMA / CTA / PV | Spell the letters |', '| Split-K | “split kay” |', '| cp.async | “see-pee ay-sync” |', '| RTX 5090 | “R–T–X fifty-ninety” |', '',
        '## Optional cuts after rehearsal', '',
        '- Slide 3: omit the final sentence about coalesced sums; its pass name and diagram remain on screen.',
        '- Slide 5: omit the final sentence about separate FP16 storage; keep the separate strip visible.',
        '- Slide 6: omit the fragment-lookahead sentence if necessary; the diagram still shows this lowering.', '',
        'Keep the tensor approximation, the original-grid qualification, LLVM attribution, and benchmark qualifiers. If the take runs long, use a shorter sentence or pickup rather than speeding up technical names.', '',
        '## Evidence and boundaries', '',
        '- The [round-6 comparison](../../benchmarks/rtx5090-2026-10-24-r6.json) has 270 balanced warm calls over 15 open workloads. Geometric means are 3.752 µs Deuteron, 6.216 µs generated CUDA/NVCC, and 4.332 µs tested Triton. These are whole-compiler results, with a limited Triton search and original checks. Compilation, assembly, and transport are excluded.',
        '- The split-K illustration shows four partitions per output tile within a larger grid. Factor selection preserves enough total parallel work, bounded K depth, and scratch capacity. The physical grid and atomic output updates remain.',
        '- The packed BF16 rewrite replaces an FP32 copy ring. Eligible TF32x3 and FP32 paths retain asynchronous copies. “One used shared stage” is not a claim that reserved shared memory halves.',
        '- BF16/TF32 are internal approximate tensor formats, with FP32 buffers and accumulation. FP16 storage is a separate pointwise subset with explicit FP32 arithmetic. Reductions can change association.',
        '- Recurrence weights are unnormalized. Each weight directly updates register mass and is also stored for PV. Shared-storage aliasing needs same-warp/component dependence; grouped PV restores the denominator for each suffix replay.',
        '- The [deployment contract](../production.md) describes the supported scheduled-IR subset. TMA, WGMMA, TMEM, public BF16 buffers, and native half arithmetic are outside the implementation.', '',
        'The previous standalone PTX-demo and reliability/test-count slides are removed from this presentation. The optional captured demo remains in the downloadable source kit as supporting material, outside the narrated deck.', '',
        'At **2:55**, stop speaking and hold slide 9 unchanged until **3:00**. The [recording and editing guide](recording-and-editing.md) describes the planned postproduction.', '',
    ]
    (destination / 'presentation-script.md').write_text('\n'.join(lines))


def render_deck(slides, path, frame_dir):
    frame_dir.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=(WIDTH / 2, HEIGHT / 2), pageCompression=1)
    pdf.setTitle('Deuteron — pass reference' if slides[0].reference else 'Deuteron — compiler passes and innovations')
    pdf.setAuthor('Deuteron')
    for slide in slides:
        draw_pdf_slide(pdf, slide)
        pdf.showPage()
    pdf.save()
    with fitz.open(path) as document:
        assert len(document) == len(slides)
        for index, page in enumerate(document):
            frame = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            assert (frame.width, frame.height) == (WIDTH, HEIGHT)
            frame.save(frame_dir / f'slide-{index + 1:02d}.png')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=HERE)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    for face, path in FONT_FILES.items():
        if not path.is_file():
            raise SystemExit(f'Missing build font {path}; update FONT_FILES for your system.')
        pdfmetrics.registerFont(TTFont(face, str(path)))
    pdfmetrics.registerFontFamily('regular', normal='regular', bold='bold', italic='regular', boldItalic='bold')
    script = read(HERE / 'script.json')
    inventory = read(HERE / 'passes.json')
    warm = read(REPO / 'benchmarks/rtx5090-2026-10-24-r6.json')
    expected = [(0,12), (12,27), (27,44), (44,63), (63,86), (86,109), (109,132), (132,150), (150,175)]
    assert [(segment['start_seconds'], segment['end_seconds']) for segment in script['segments']] == expected
    slides = create_slides(warm)
    references = create_reference_slides(inventory)
    frames, reference_frames = output / 'slides', output / 'reference-slides'
    render_deck(slides, output / 'slides.pdf', frames)
    render_deck(references, output / 'pass-reference.pdf', reference_frames)
    build_pptx(slides, script, output / 'slides.pptx', inventory)
    build_pptx(references, None, output / 'pass-reference.pptx', inventory)
    word_count = make_script_documents(script, output)
    make_presentation_guide(script, inventory, output)
    previews = Image.new('RGB', (1920, 1280), '#060C16')
    for index in range(len(slides)):
        with Image.open(frames / f'slide-{index + 1:02d}.png') as original:
            thumbnail = original.resize((592, 333), Image.Resampling.LANCZOS)
            previews.paste(thumbnail, (28 + (index % 3) * 636, 128 + (index // 3) * 382))
    draw = ImageDraw.Draw(previews)
    draw.text((28, 23), 'DEUTERON / COMPILER PASSES AND INNOVATIONS', font=font(33, 'bold'), fill=INK)
    draw.text((30, 78), f'9 narrated slides · {word_count} words · 3 separate reference slides · 3 minutes', font=font(24), fill=MUTED)
    for index, segment in enumerate(script['segments']):
        stop = segment['end_seconds'] if index < len(slides)-1 else 180
        label = f"{index + 1:02d}  {timestamp(segment['start_seconds'])}–{timestamp(stop)}"
        draw.text((38 + (index % 3) * 636, 465 + (index // 3) * 382), label, font=font(22, 'bold'), fill=MINT)
    previews.save(output / 'slides-overview.png')
    source_files = [HERE / name for name in ('script.json', 'passes.json', 'build_assets.py', 'requirements.txt', 'innovations.md', 'pass-reference.md')]
    source_files += [REPO / 'README.md', REPO / 'benchmarks/rtx5090-2026-10-24-r6.json']
    source_files += sorted({REPO / entry['source'] for entry in inventory['entries']})
    artifacts = [output / name for name in ('slides.pdf', 'slides.pptx', 'pass-reference.pdf', 'pass-reference.pptx',
                                          'recording-script.pdf', 'recording-script.docx', 'narration.txt', 'slides-overview.png', 'presentation-script.md')]
    artifacts += [frames / f'slide-{i:02d}.png' for i in range(1,10)]
    artifacts += [reference_frames / f'slide-{i:02d}.png' for i in range(1,4)]
    manifest = {'generated_utc': datetime.now(timezone.utc).isoformat(), 'revision': 'passes-and-innovations',
                'source_release': 'efb728d9106d864bb5f35efc95bfb5f01d95646a', 'spoken_words': word_count,
                'timeline_seconds': 180, 'narration_ends_seconds': 175, 'slides': len(slides),
                'reference_slides': len(references), 'reference_slides_in_video': False, 'pass_entries': 27,
                'frame_size': [WIDTH, HEIGHT], 'overview_size': [1920, 1280], 'final_video_status': 'awaiting_user_recording',
                'source_sha256': {str(path.relative_to(REPO)): sha(path) for path in source_files},
                'artifacts': {str(path.relative_to(output)): {'sha256': sha(path), 'bytes': path.stat().st_size} for path in artifacts},
                'font_files': {key: {'path': str(value), 'sha256': sha(value)} for key, value in FONT_FILES.items()},
                'layout': {'all_text_measured_with_build_fonts': True, 'all_text_within_canvas': True,
                           'pdf_and_png_share_rendering': True, 'pptx_uses_editable_text_and_shapes': True,
                           'pptx_rendering_note': 'PowerPoint rendering depends on installed fonts. PDF/PNG are the fixed visual exports.'}}
    (output / 'assets-manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps({'output': str(output), 'slides': len(slides), 'reference_slides': len(references),
                      'spoken_words': word_count, 'files': len(artifacts), 'timeline_seconds': 180}, indent=2))


if __name__ == '__main__':
    main()
